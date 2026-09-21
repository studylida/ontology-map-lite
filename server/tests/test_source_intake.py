from io import BytesIO

import pytest
from docx import Document

from ontology_map.source_intake import (
    SourceIntakeError,
    extract_normalized_body,
    parse_published_at,
)


def test_text_upload_uses_existing_normalization_contract() -> None:
    body = extract_normalized_body(
        "  첫 줄  \r\n\r\n\r\n둘째 줄  \n".encode(),
        filename="article.txt",
        content_type="text/plain",
    )
    assert body == "첫 줄\n\n둘째 줄"


def test_docx_upload_extracts_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("첫 문단")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "왼쪽"
    table.rows[0].cells[1].text = "오른쪽"

    body = extract_normalized_body(
        _save_docx(document),
        filename="article.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert body == "첫 문단\n왼쪽\t오른쪽"


def _save_docx(document: Document) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_upload_limits_and_rejects_unknown_format() -> None:
    with pytest.raises(SourceIntakeError, match="10 MiB") as too_large:
        extract_normalized_body(
            b"x" * (10 * 1024 * 1024 + 1), filename="large.txt", content_type=None
        )
    assert too_large.value.code == "SOURCE_TOO_LARGE"

    with pytest.raises(SourceIntakeError) as unsupported:
        extract_normalized_body(b"body", filename="article.rtf", content_type=None)
    assert unsupported.value.code == "SOURCE_FORMAT_UNSUPPORTED"


def test_published_date_is_anchored_at_utc_midnight() -> None:
    assert parse_published_at("2026-09-20").isoformat() == "2026-09-20T00:00:00+00:00"
    assert parse_published_at(None) is None
