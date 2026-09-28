# server/tests/test_document_parser.py
"""문서 파서 단위 테스트."""

import io
import docx
from ontology_map.services.document_parser import (
    clean_html_to_text,
    parse_docx,
    extract_document_text,
)


def test_clean_html():
    """HTML 태그와 스크립트가 안전하게 제거되는지 검증."""
    html = "<html><head><script>alert('x');</script></head><body><h1>지식 온톨로지</h1><p>본문 내용입니다.</p></body></html>"
    text = clean_html_to_text(html)
    assert "alert" not in text
    assert "지식 온톨로지" in text
    assert "본문 내용입니다." in text


def test_parse_docx():
    """DOCX 워드 파일 바이너리에서 단락 텍스트 추출 검증."""
    doc = docx.Document()
    doc.add_heading("테스트 공고문", level=1)
    doc.add_paragraph("중소벤처기업부 지원사업 본문 내용입니다.")
    
    bio = io.BytesIO()
    doc.save(bio)
    file_bytes = bio.getvalue()

    extracted = parse_docx(file_bytes)
    assert "테스트 공고문" in extracted
    assert "중소벤처기업부" in extracted


def test_extract_document_text_text_mode():
    """텍스트 모드 및 파일명 추출 정규화 검증."""
    title, content = extract_document_text("text", "인공지능 연구소 메모", "AI_Memo.txt")
    assert title == "AI_Memo.txt"
    assert content == "인공지능 연구소 메모"
